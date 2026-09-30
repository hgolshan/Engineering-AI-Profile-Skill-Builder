#!/usr/bin/env python3
"""
Engineering AI Profile & Skill Builder - Word Document Proposal Generator
Optional local utility to convert docs/PROJECT_PROPOSAL.md into a formatted Word (.docx) proposal.

Requirements:
    pip install python-docx

Usage:
    python3 docs/build_docx.py
"""

import sys
import os
import re

def main():
    try:
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
        from docx.oxml import OxmlElement, parse_xml
        from docx.oxml.ns import nsdecls, qn
    except ImportError:
        print("\n" + "=" * 70)
        print("ERROR: 'python-docx' is not installed.")
        print("To generate the Word document, please install the package:")
        print("    pip install python-docx")
        print("=" * 70 + "\n")
        sys.exit(1)

    # Determine paths
    script_dir = os.path.dirname(os.path.abspath(__file__))
    md_path = os.path.join(script_dir, "PROJECT_PROPOSAL.md")
    output_docx = os.path.join(script_dir, "Engineering_AI_Profile_Skill_Builder_Proposal.docx")

    if not os.path.exists(md_path):
        print(f"Error: Could not find proposal markdown at {md_path}")
        sys.exit(1)

    print(f"Reading markdown proposal: {md_path}")
    with open(md_path, "r", encoding="utf-8") as f:
        md_content = f.read()

    # Create Word document
    doc = Document()

    # Configure Margins (1 inch all around)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Style colors
    PRIMARY_COLOR = RGBColor(15, 23, 42)      # Dark Slate
    ACCENT_COLOR = RGBColor(37, 99, 235)      # Engineering Blue
    MUTED_COLOR = RGBColor(100, 116, 139)     # Slate Gray

    # Helper function for setting cell shading
    def set_cell_background(cell, hex_color):
        shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
        cell._tc.get_or_add_tcPr().append(shading_elm)

    # Helper function for setting cell margins
    def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = OxmlElement('w:tcMar')
        for margin_name, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
            node = OxmlElement(margin_name)
            node.set(qn('w:w'), str(val))
            node.set(qn('w:type'), 'dxa')
            tcMar.append(node)
        tcPr.append(tcMar)

    # Document Header Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("Engineering AI Profile & Skill Builder")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(24)
    run_title.font.bold = True
    run_title.font.color.rgb = PRIMARY_COLOR

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(16)
    run_sub = sub_p.add_run("Formal Project Proposal & Strategic Implementation Framework")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(14)
    run_sub.font.bold = True
    run_sub.font.color.rgb = ACCENT_COLOR

    # Metadata Box
    meta_table = doc.add_table(rows=5, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    meta_items = [
        ("Document Identifier", "PRJ-ENG-AI-2026-001"),
        ("Project Lead & Author", "Hossein Golshan (github.com/hgolshan)"),
        ("Target Industry", "Heavy Industry, Metal Smelters, Mineral Processing & EPC Projects"),
        ("Production Hosting", "GitHub Pages (https://hgolshan.github.io/Engineering-AI-Profile-Skill-Builder/)"),
        ("License & Governance", "Open-Source under MIT License")
    ]

    for i, (label, val) in enumerate(meta_items):
        row = meta_table.rows[i]
        cell_lbl, cell_val = row.cells[0], row.cells[1]
        cell_lbl.width = Inches(2.2)
        cell_val.width = Inches(4.3)
        set_cell_background(cell_lbl, "F1F5F9")
        set_cell_background(cell_val, "F8FAFC")
        set_cell_margins(cell_lbl, top=80, bottom=80, left=120, right=120)
        set_cell_margins(cell_val, top=80, bottom=80, left=120, right=120)

        p0 = cell_lbl.paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(label)
        r0.font.name = "Arial"
        r0.font.size = Pt(9.5)
        r0.font.bold = True
        r0.font.color.rgb = PRIMARY_COLOR

        p1 = cell_val.paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(val)
        r1.font.name = "Arial"
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = PRIMARY_COLOR

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Process markdown lines
    lines = md_content.split("\n")
    in_table = False
    table_lines = []

    def flush_table(t_lines):
        if not t_lines:
            return
        parsed_rows = []
        for line in t_lines:
            # strip start/end pipe
            stripped = line.strip()
            if not stripped.startswith("|"):
                continue
            cells = [c.strip() for c in stripped.split("|")[1:-1]]
            # check if separator row
            if all(re.match(r"^:?-+:?$", c) for c in cells):
                continue
            parsed_rows.append(cells)

        if not parsed_rows:
            return

        col_count = max(len(r) for r in parsed_rows)
        tbl = doc.add_table(rows=len(parsed_rows), cols=col_count)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = True

        for r_idx, row_data in enumerate(parsed_rows):
            is_header = (r_idx == 0)
            for c_idx in range(col_count):
                cell_text = row_data[c_idx] if c_idx < len(row_data) else ""
                cell = tbl.rows[r_idx].cells[c_idx]
                set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
                if is_header:
                    set_cell_background(cell, "1E293B")
                else:
                    set_cell_background(cell, "FFFFFF" if r_idx % 2 == 1 else "F8FAFC")

                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                r = p.add_run(cell_text)
                r.font.name = "Arial"
                r.font.size = Pt(9)
                if is_header:
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(255, 255, 255)
                else:
                    r.font.color.rgb = PRIMARY_COLOR

        p_after = doc.add_paragraph()
        p_after.paragraph_format.space_after = Pt(10)

    # Line by line processing
    for line in lines:
        stripped = line.strip()

        # Handle tables
        if stripped.startswith("|"):
            in_table = True
            table_lines.append(stripped)
            continue
        elif in_table:
            in_table = False
            flush_table(table_lines)
            table_lines = []

        # Skip main title and metadata since we formatted them custom
        if stripped.startswith("# Project Proposal:") or stripped.startswith("**Document Identifier:**") or \
           stripped.startswith("**Project Title:**") or stripped.startswith("**Project Author") or \
           stripped.startswith("**Target Organization:**") or stripped.startswith("**Distribution:**") or \
           stripped.startswith("**Production URL:**") or stripped.startswith("**Source Repository:**"):
            continue

        if stripped == "---":
            continue

        # Headings
        if stripped.startswith("## "):
            h = doc.add_heading(level=1)
            h.paragraph_format.space_before = Pt(16)
            h.paragraph_format.space_after = Pt(6)
            r = h.add_run(stripped[3:].strip())
            r.font.name = "Arial"
            r.font.size = Pt(14)
            r.font.bold = True
            r.font.color.rgb = ACCENT_COLOR
        elif stripped.startswith("### "):
            h = doc.add_heading(level=2)
            h.paragraph_format.space_before = Pt(12)
            h.paragraph_format.space_after = Pt(4)
            r = h.add_run(stripped[4:].strip())
            r.font.name = "Arial"
            r.font.size = Pt(11.5)
            r.font.bold = True
            r.font.color.rgb = PRIMARY_COLOR
        elif stripped.startswith("#### "):
            h = doc.add_heading(level=3)
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(2)
            r = h.add_run(stripped[5:].strip())
            r.font.name = "Arial"
            r.font.size = Pt(10.5)
            r.font.bold = True
            r.font.color.rgb = MUTED_COLOR
        elif stripped.startswith("> "):
            # Blockquote / Callout box
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.right_indent = Inches(0.4)
            r = p.add_run(stripped[2:].strip().replace("*", ""))
            r.font.name = "Arial"
            r.font.size = Pt(9.5)
            r.font.italic = True
            r.font.color.rgb = RGBColor(180, 83, 9) # Amber warning
        elif stripped.startswith("- ") or stripped.startswith("* "):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            # Remove bold markup for clean rendering
            clean_text = stripped[2:].strip()
            # Split bold markers
            parts = re.split(r'(\*\*.*?\*\*)', clean_text)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                else:
                    r = p.add_run(part)
                r.font.name = "Arial"
                r.font.size = Pt(10)
                r.font.color.rgb = PRIMARY_COLOR
        elif re.match(r"^\d+\.\s+", stripped):
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            match = re.match(r"^\d+\.\s+(.*)$", stripped)
            clean_text = match.group(1) if match else stripped
            parts = re.split(r'(\*\*.*?\*\*)', clean_text)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                else:
                    r = p.add_run(part)
                r.font.name = "Arial"
                r.font.size = Pt(10)
                r.font.color.rgb = PRIMARY_COLOR
        elif stripped:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(6)
            parts = re.split(r'(\*\*.*?\*\*)', stripped)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                else:
                    r = p.add_run(part)
                r.font.name = "Arial"
                r.font.size = Pt(10)
                r.font.color.rgb = PRIMARY_COLOR

    # Flush any remaining table
    if in_table and table_lines:
        flush_table(table_lines)

    # Save document
    doc.save(output_docx)
    print(f"\nSUCCESS: Generated Word proposal document successfully at:")
    print(f"  --> {output_docx}\n")

if __name__ == "__main__":
    main()
